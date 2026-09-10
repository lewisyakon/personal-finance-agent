from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.parsers.wechat_csv import _header_index, parse_wechat_rows
from app.schemas.transaction import ParseReport


def _worksheet_rows(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            if _header_index(rows) is not None:
                return rows
        if workbook.worksheets:
            return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
        return []
    finally:
        workbook.close()


class WechatXlsxParser:
    source = "wechat"
    format = "xlsx"

    def parse(self, path: Path) -> ParseReport:
        try:
            rows = _worksheet_rows(path)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise ValueError("XLSX 文件无法读取，请确认文件未损坏") from exc
        return parse_wechat_rows(rows, encoding="xlsx", format=self.format)
